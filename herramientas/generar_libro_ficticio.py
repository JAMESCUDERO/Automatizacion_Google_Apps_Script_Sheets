"""Genera el libro de ejemplo FICTICIO que usa el script de Apps Script.

TODOS LOS DATOS SON FICTICIOS: razones sociales ("EMPRESA FICTICIA 0001 S.A."), RUC con prefijo 99
(provincia inexistente), documentos "DOC-FICT-..." y funcionarios "FUNCIONARIO A/B/C/D".

Estructura exigida por apps-script/Codigo.gs (ver CONFIG):
  hoja SOLICITUDES  - fila 1 encabezados; columna C = tipo de solicitud; columna I = funcionario asignado
                      (STATUS_COLUMN = 9); columna J = fecha de modificación (timestamp automático)
  hoja LISTAS       - N21 = tipo consultado (lo escribe el script); O21 = funcionario sugerido (fórmula)

Uso:  python herramientas/generar_libro_ficticio.py
"""
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parents[1]
DESTINO = RAIZ / "datos-ejemplo" / "SOLICITUDES_FICTICIAS.xlsx"
N_FILAS = 900                       # CONTEXT_MIN_ROW = 820: las filas > 820 son "solicitudes nuevas"
SEMILLA = 20240601

TIPOS = ["01-AUTORIZACION", "02-RENOVACION", "03-CAMBIO DE DOMICILIO", "04-MODIFICACION",
         "05-CANCELACION", "06-CONSULTA", "07-OTROS"]
PESOS_TIPO = [0.22, 0.28, 0.16, 0.12, 0.06, 0.10, 0.06]
OPERADORES = ["01-IMPORTADOR", "02-EXPORTADOR", "03-TRANSPORTISTA", "04-AGENTE", "05-DEPOSITO", "06-COURIER"]
FUNCIONARIOS = ["01-FUNCIONARIO A", "02-FUNCIONARIO B", "03-FUNCIONARIO C", "04-FUNCIONARIO D"]
ESTADOS = ["0-ASIGNADO", "1-EN REVISION", "2-INSPECCION", "3-COMPLETAR", "4-RESUELTO", "5-ARCHIVADO"]

ENCABEZADOS = ["N", "RAZON SOCIAL", "TIPO DE SOLICITUD", "RUC", "TIPO DE OPERADOR", "DOCUMENTO DE INGRESO",
               "FECHA DE SOLICITUD", "FECHA DE ADMISION", "FUNCIONARIO ASIGNADO", "FECHA DE MODIFICACION",
               "ESTADO TRAMITE"]
AZUL = PatternFill("solid", fgColor="1F3864")
GRIS = PatternFill("solid", fgColor="D9E1F2")


def hoja_solicitudes(wb, rng):
    ws = wb.active
    ws.title = "SOLICITUDES"
    ws.append(ENCABEZADOS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), AZUL
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    tipos = rng.choice(TIPOS, N_FILAS, p=PESOS_TIPO)
    base = date(2024, 1, 2)
    for i in range(N_FILAS):
        n = i + 1
        solicitud = base + timedelta(days=int(rng.integers(0, 355)))
        admision = solicitud + timedelta(days=int(rng.integers(1, 6)))
        pendiente = n > N_FILAS - 20                       # últimas 20: aún sin funcionario asignado
        funcionario = None if pendiente else FUNCIONARIOS[int(rng.integers(0, len(FUNCIONARIOS)))]
        modificacion = None if pendiente else datetime.combine(admision, datetime.min.time()) + timedelta(
            days=int(rng.integers(0, 10)), hours=int(rng.integers(8, 18)))
        estado = ESTADOS[0] if pendiente else ESTADOS[int(rng.integers(0, len(ESTADOS)))]
        ws.append([
            n, f"EMPRESA FICTICIA {int(rng.integers(1, 481)):04d} S.A.", str(tipos[i]),
            f"99{int(rng.integers(0, 10**8)):08d}001", OPERADORES[int(rng.integers(0, len(OPERADORES)))],
            f"DOC-FICT-2024-{n:06d}", solicitud, admision, funcionario, modificacion, estado,
        ])
    for fila in ws.iter_rows(min_row=2, min_col=7, max_col=8):
        for c in fila:
            c.number_format = "dd/mm/yyyy"
    for fila in ws.iter_rows(min_row=2, min_col=10, max_col=10):
        for c in fila:
            c.number_format = "dd/mm/yyyy hh:mm"
    for col, ancho in zip("ABCDEFGHIJK", [6, 30, 26, 16, 20, 24, 14, 14, 22, 18, 18]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"

    # Listas desplegables como en el modelo original (apuntan a LISTAS)
    for rango, formula in [("C2:C2000", "=LISTAS!$A$3:$A$9"), ("E2:E2000", "=LISTAS!$C$3:$C$8"),
                           ("I2:I2000", "=LISTAS!$D$3:$D$6"), ("K2:K2000", "=LISTAS!$B$3:$B$8")]:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(rango)


def hoja_listas(wb):
    ws = wb.create_sheet("LISTAS")
    ws["A1"] = "LISTAS DE APOYO Y MOTOR DE REGLAS - DATOS FICTICIOS"
    ws["A1"].font = Font(bold=True, size=12)
    for col, titulo, valores in [("A", "TIPO DE SOLICITUD", TIPOS), ("B", "ESTADO DE TRAMITE", ESTADOS),
                                 ("C", "TIPO DE OPERADOR", OPERADORES), ("D", "FUNCIONARIO ASIGNADO", FUNCIONARIOS)]:
        ws[f"{col}2"] = titulo
        ws[f"{col}2"].font, ws[f"{col}2"].fill = Font(bold=True), GRIS
        for k, v in enumerate(valores):
            ws[f"{col}{3 + k}"] = v
        ws.column_dimensions[col].width = 26

    # Carga de trabajo por tipo y funcionario (conteo vivo sobre SOLICITUDES)
    ws["N2"] = "CARGA ACTUAL (conteo de solicitudes)"
    ws["N2"].font = Font(bold=True)
    ws["N3"] = "TIPO DE SOLICITUD"
    for j, f in enumerate(FUNCIONARIOS):
        ws.cell(row=3, column=15 + j, value=f)
    for c in ws[3][13:19]:
        c.font, c.fill = Font(bold=True), GRIS
    for i, tipo in enumerate(TIPOS):
        fila = 4 + i
        ws.cell(row=fila, column=14, value=tipo)
        for j in range(len(FUNCIONARIOS)):
            col_letra = ws.cell(row=3, column=15 + j).column_letter
            ws.cell(row=fila, column=15 + j,
                    value=f"=COUNTIFS(SOLICITUDES!$C:$C,$N{fila},SOLICITUDES!$I:$I,{col_letra}$3)")
    ws.column_dimensions["N"].width = 26
    for col in "OPQR":
        ws.column_dimensions[col].width = 20

    # Motor de reglas: N21 lo escribe el script (onEdit); O21 sugiere al funcionario con menor carga
    ws["N20"], ws["O20"] = "TIPO DE SOLICITUD CONSULTADA", "FUNCIONARIO SUGERIDO"
    for c in (ws["N20"], ws["O20"]):
        c.font, c.fill = Font(bold=True), GRIS
    ws["N21"] = TIPOS[1]
    ws["O21"] = ('=IFERROR(INDEX($O$3:$R$3,MATCH(MIN(INDEX($O$4:$R$10,MATCH($N$21,$N$4:$N$10,0),0)),'
                 'INDEX($O$4:$R$10,MATCH($N$21,$N$4:$N$10,0),0),0)),"")')
    ws["N23"] = ("Regla de ejemplo: se sugiere al funcionario con MENOR carga para el tipo consultado. "
                 "Las reglas reales se centralizan aquí, sin tocar el script.")
    ws["N23"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("N23:R25")


def hoja_leeme(wb):
    ws = wb.create_sheet("LEEME", 0)
    lineas = [
        ("DATOS 100 % FICTICIOS", True),
        ("Este libro es una muestra para probar el script de Apps Script. Ninguna empresa, RUC, documento ni "
         "funcionario es real.", False),
        ("", False),
        ("Convenciones: razones sociales 'EMPRESA FICTICIA ####', RUC con prefijo 99, documentos 'DOC-FICT-...', "
         "funcionarios 'FUNCIONARIO A-D'.", False),
        ("", False),
        ("Uso: importar en Google Sheets (Archivo > Importar), abrir Extensiones > Apps Script y pegar "
         "apps-script/Codigo.gs.", False),
        ("Prueba: edite la columna I (funcionario) de SOLICITUDES y vea la marca de tiempo en la columna J; "
         "edite la columna C en una fila posterior a la 820 y vea LISTAS!N21 y O21.", False),
    ]
    for i, (texto, negrita) in enumerate(lineas, start=1):
        c = ws.cell(row=i, column=1, value=texto)
        c.font = Font(bold=negrita, size=14 if negrita else 11, color="C00000" if negrita else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110


def main():
    rng = np.random.default_rng(SEMILLA)
    wb = Workbook()
    hoja_solicitudes(wb, rng)
    hoja_listas(wb)
    hoja_leeme(wb)
    DESTINO.parent.mkdir(exist_ok=True)
    wb.save(DESTINO)
    print(f"Generado: {DESTINO} ({DESTINO.stat().st_size / 1024:.0f} KB) | solicitudes: {N_FILAS}")


if __name__ == "__main__":
    main()
